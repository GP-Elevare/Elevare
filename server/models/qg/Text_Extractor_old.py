from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

def extract_text_from_pptx(pptx_path: str, include_notes: bool = False) -> str:
    """
    Extract all text content from a .pptx file.

    Args:
        pptx_path (str): Path to the .pptx file.
        include_notes (bool): If True, also extracts speaker notes.

    Returns:
        str: Extracted text as a single string (slides separated clearly).
    """
    prs = Presentation(pptx_path)
    output_lines = []

    for slide_idx, slide in enumerate(prs.slides, start=1):
        slide_lines = [f"--- Slide {slide_idx} ---"]

        # Extract text from shapes (text boxes, placeholders, etc.)
        for shape in slide.shapes:
            # Regular text frames
            if hasattr(shape, "text_frame") and shape.has_text_frame:
                text = shape.text.strip()
                if text:
                    slide_lines.append(text)

            # Tables
            if shape.shape_type == MSO_SHAPE_TYPE.TABLE:
                table = shape.table
                for row in table.rows:
                    row_text = []
                    for cell in row.cells:
                        cell_text = cell.text.strip()
                        if cell_text:
                            row_text.append(cell_text)
                    if row_text:
                        slide_lines.append(" | ".join(row_text))

        # Extract speaker notes (optional)
        if include_notes:
            try:
                notes_slide = slide.notes_slide
                if notes_slide and notes_slide.notes_text_frame:
                    notes_text = notes_slide.notes_text_frame.text.strip()
                    if notes_text:
                        slide_lines.append("[Notes]")
                        slide_lines.append(notes_text)
            except Exception:
                # Some slides may not have notes
                pass

        output_lines.append("\n".join(slide_lines))

    return "\n\n".join(output_lines)


# Example usage:
# text = extract_text_from_pptx("my_presentation.pptx", include_notes=True)
# print(text)
